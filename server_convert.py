#!/usr/bin/env python3
"""
AirSlide — PPTX Conversion Server
Accepts a .pptx upload, renders each slide to a high-res image using
python-pptx + Pillow, then returns a multi-page PDF for the existing
PDFRenderer on the client.

Usage:
    python3 server_convert.py          # starts on port 5001
    python3 server_convert.py 5002     # custom port

Install deps once:
    pip3 install python-pptx Pillow flask flask-cors
"""
import io
import os
import sys

from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

# ─── App Setup ────────────────────────────────────────────────────────────────
app = Flask(__name__)
CORS(app, resources={r"/api/*": {"origins": "*"}})

# Output resolution — 1920×1080 at 16:9
SLIDE_W = 1920
SLIDE_H = 1080

# Font search paths across macOS / Linux / Windows
_FONT_PATHS = [
    "/System/Library/Fonts/Helvetica.ttc",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "/System/Library/Fonts/SFNS.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
    "C:/Windows/Fonts/arial.ttf",
]
_font_cache: dict = {}


def get_font(size: int) -> ImageFont.ImageFont:
    size = max(8, size)
    if size in _font_cache:
        return _font_cache[size]
    for path in _FONT_PATHS:
        if os.path.exists(path):
            try:
                f = ImageFont.truetype(path, size)
                _font_cache[size] = f
                return f
            except Exception:
                pass
    f = ImageFont.load_default()
    _font_cache[size] = f
    return f


# ─── Helpers ──────────────────────────────────────────────────────────────────

def rgb_tuple(rgb):
    """Convert pptx RGBColor → (r, g, b) tuple, or None on failure."""
    try:
        return (rgb.r, rgb.g, rgb.b)
    except Exception:
        return None


def scale_x(emu, slide_emu):
    return int(emu / slide_emu * SLIDE_W)


def scale_y(emu, slide_emu):
    return int(emu / slide_emu * SLIDE_H)


def get_bg_color(slide):
    """Return (r, g, b) background colour for the slide, defaulting to white."""
    try:
        fill = slide.background.fill
        if fill.type is not None:
            name = fill.type.name
            if name == "SOLID":
                c = rgb_tuple(fill.fore_color.rgb)
                if c:
                    return c
    except Exception:
        pass
    return (255, 255, 255)


def draw_wrapped_text(draw, text, x1, y1, x2, y2, font, color):
    """Word-wrap text inside the bounding box."""
    words = text.split()
    lines, line = [], ""
    for word in words:
        candidate = f"{line} {word}".strip()
        try:
            bbox = draw.textbbox((0, 0), candidate, font=font)
            if bbox[2] - bbox[0] > (x2 - x1 - 16) and line:
                lines.append(line)
                line = word
            else:
                line = candidate
        except Exception:
            line = candidate
    if line:
        lines.append(line)

    ty = y1 + 6
    try:
        line_h = draw.textbbox((0, 0), "Ay", font=font)[3] + 4
    except Exception:
        line_h = 20
    for ln in lines:
        if ty + line_h > y2:
            break
        draw.text((x1 + 8, ty), ln, fill=color, font=font)
        ty += line_h


# ─── Core Slide Renderer ──────────────────────────────────────────────────────

def render_slide(slide, prs) -> Image.Image:
    sw = prs.slide_width   # total width in EMU
    sh = prs.slide_height  # total height in EMU

    # Background
    img = Image.new("RGB", (SLIDE_W, SLIDE_H), get_bg_color(slide))
    draw = ImageDraw.Draw(img)

    for shape in slide.shapes:
        try:
            if shape.left is None or shape.top is None:
                continue

            x1 = scale_x(shape.left, sw)
            y1 = scale_y(shape.top, sh)
            w  = scale_x(shape.width, sw)
            h  = scale_y(shape.height, sh)
            x2, y2 = x1 + w, y1 + h

            if w <= 0 or h <= 0:
                continue

            # ── Pictures ──────────────────────────────────────────────────────
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = shape.image.blob
                    pic = Image.open(io.BytesIO(blob)).convert("RGBA")
                    pic = pic.resize((w, h), Image.LANCZOS)
                    r, g, b, a = pic.split()
                    img.paste(Image.merge("RGB", (r, g, b)), (x1, y1), a)
                except Exception as e:
                    print(f"  [warn] picture render: {e}", file=sys.stderr)

            # ── Text Frames ───────────────────────────────────────────────────
            elif shape.has_text_frame:
                # Shape fill background
                try:
                    fill = shape.fill
                    if fill.type is not None and fill.type.name == "SOLID":
                        fc = rgb_tuple(fill.fore_color.rgb)
                        if fc:
                            draw.rectangle([x1, y1, x2, y2], fill=fc)
                except Exception:
                    pass

                # Render paragraphs
                ty = y1 + 6
                for para in shape.text_frame.paragraphs:
                    line = para.text.strip()
                    if not line:
                        ty += 8
                        continue

                    # Style from first run
                    font_pt = 18
                    color = (0, 0, 0)
                    for run in para.runs:
                        try:
                            if run.font.size and run.font.size.pt:
                                font_pt = max(6, int(run.font.size.pt))
                            c = rgb_tuple(run.font.color.rgb)
                            if c:
                                color = c
                        except Exception:
                            pass
                        break

                    px_size = max(8, int(font_pt * SLIDE_H / 720))
                    font = get_font(px_size)
                    draw_wrapped_text(draw, line, x1, ty, x2, y2, font, color)
                    try:
                        line_h = draw.textbbox((0, 0), "Ay", font=font)[3] + 6
                    except Exception:
                        line_h = px_size + 6
                    ty += line_h

            # ── Solid shapes (rectangles, etc.) ───────────────────────────────
            else:
                try:
                    fill = shape.fill
                    if fill.type is not None and fill.type.name == "SOLID":
                        fc = rgb_tuple(fill.fore_color.rgb)
                        if fc:
                            draw.rectangle([x1, y1, x2, y2], fill=fc)
                except Exception:
                    pass

        except Exception as e:
            print(f"  [warn] shape: {e}", file=sys.stderr)

    return img


# ─── Routes ───────────────────────────────────────────────────────────────────

@app.route("/api/health", methods=["GET"])
def health():
    return jsonify({"status": "ok", "service": "AirSlide PPTX Converter v1"})


@app.route("/api/convert", methods=["POST"])
def convert_pptx():
    if "file" not in request.files:
        return jsonify({"error": "No file field in request"}), 400

    upload = request.files["file"]
    if not upload.filename.lower().endswith(".pptx"):
        return jsonify({"error": "Only .pptx files are accepted"}), 400

    try:
        raw = upload.read()
        prs = Presentation(io.BytesIO(raw))
        total = len(prs.slides)

        if total == 0:
            return jsonify({"error": "Presentation has no slides"}), 400

        print(f"  Converting {upload.filename!r} ({total} slides)…", file=sys.stderr)

        images = []
        for i, slide in enumerate(prs.slides):
            print(f"    slide {i + 1}/{total}", file=sys.stderr)
            images.append(render_slide(slide, prs))

        # Save to PDF in-memory
        buf = io.BytesIO()
        if len(images) == 1:
            images[0].save(buf, format="PDF", resolution=150)
        else:
            images[0].save(
                buf,
                format="PDF",
                save_all=True,
                append_images=images[1:],
                resolution=150,
            )
        buf.seek(0)

        print(f"  ✓ Done — {buf.getbuffer().nbytes // 1024} KB PDF", file=sys.stderr)

        return send_file(
            buf,
            mimetype="application/pdf",
            as_attachment=False,
            download_name="converted.pdf",
        )

    except Exception as e:
        import traceback
        traceback.print_exc(file=sys.stderr)
        return jsonify({"error": str(e)}), 500


# ─── Entry Point ──────────────────────────────────────────────────────────────

if __name__ == "__main__":
    port = int(sys.argv[1]) if len(sys.argv) > 1 else 5001
    print(f"\n✋  AirSlide PPTX Converter  →  http://localhost:{port}")
    print("   POST /api/convert   multipart: file=<.pptx>  →  PDF")
    print("   GET  /api/health\n")
    app.run(host="0.0.0.0", port=port, debug=False)
